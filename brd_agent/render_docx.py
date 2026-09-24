"""Render the BRD and SRS into formatted Word documents.

Narrative sections come from the BRD/SRS models. Requirement, conflict, open-question
and traceability tables come straight from the validated registry (`ExtractedFacts`),
so nothing reaches the documents without passing validation.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from .audit import summary
from .models import BRD, SRS, ExtractedFacts, Requirement, Risk, SourceDoc, Stakeholder

HEADER_FILL = "1F3864"
NOT_CONFIRMED_FILL = "FFF2CC"
CONFLICT_FILL = "F8CBAD"
DRAFT_STATUS = "Draft for Review"
EMPTY = "TBD — not covered by the provided sources."


@dataclass
class RunInfo:
    """What the documents report about this generation run."""
    version: str = "v0.1"
    review_notes: list[str] = field(default_factory=list)
    change_log: dict[str, list[str]] = field(default_factory=dict)


# ---------------------------------------------------------------- helpers

def _new_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    return doc


def _italic(doc: Document, text: str) -> None:
    doc.add_paragraph().add_run(text).italic = True


def _title_page(doc: Document, short: str, title: str, project: str,
                facts: ExtractedFacts, source_count: int, run: RunInfo) -> None:
    for _ in range(4):
        doc.add_paragraph()
    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = heading.add_run(f"{short} — {DRAFT_STATUS}")
    r.bold = True
    r.font.size = Pt(28)
    r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(f"{title}\n{project}").font.size = Pt(16)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Status: {DRAFT_STATUS}").bold = True
    meta.add_run(f"\n{short} {run.version} · generated {date.today().isoformat()}")
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run(
        "Generated automatically from project sources. Nothing in this document is approved. "
        "Reviewers must resolve all conflicts, TBDs, assumptions, open questions and "
        "low-confidence items before it is considered final."
    ).italic = True

    doc.add_paragraph()
    _table(doc, ["Generation summary", "Count"],
           [[k, str(v)] for k, v in summary(facts, source_count).items()])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def _table(doc: Document, headers: list[str], rows: list[list[str]],
           fills: list[str | None] | None = None) -> None:
    if not rows:
        _italic(doc, EMPTY)
        return
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, HEADER_FILL)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
            if fills and fills[i]:
                _shade(cell, fills[i])
    doc.add_paragraph()


def _bullets(doc: Document, items: list[str]) -> None:
    if not items:
        _italic(doc, EMPTY)
        return
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _para(doc: Document, text: str) -> None:
    if not (text or "").strip():
        _italic(doc, EMPTY)
        return
    for block in text.split("\n\n"):
        doc.add_paragraph(block.strip())


def _ids(ids: list[str]) -> str:
    return ", ".join(ids) if ids else "-"


def _fill(status: str) -> str | None:
    if status == "CONFLICT":
        return CONFLICT_FILL
    return None if status == "CONFIRMED" else NOT_CONFIRMED_FILL


def _legend(doc: Document) -> None:
    _italic(doc, "Only CONFIRMED items are backed by explicit source evidence. Highlighted rows are "
                 "ASSUMPTION, TBD, NEEDS_CLARIFICATION or CONFLICT and must be reviewed.")


def _requirement_table(doc: Document, reqs: list[Requirement], with_category: bool = False,
                       with_links: bool = False) -> None:
    headers = (["Id"] + (["Category"] if with_category else [])
               + ["Requirement", "Status", "Confidence", "Priority"]
               + (["Implements"] if with_links else []) + ["Sources", "Evidence"])
    rows = []
    for r in reqs:
        text = f"{r.title}: {r.requirement}" if r.title else r.requirement
        if r.history:
            text += "\nHistory: " + "; ".join(r.history)
        rows.append([r.id] + ([r.category] if with_category else [])
                    + [text, r.status, r.confidence, r.priority]
                    + ([_ids(r.br_refs)] if with_links else [])
                    + [_ids(r.source_ids), r.source_reference or "-"])
    _table(doc, headers, rows, [_fill(r.status) for r in reqs])


def _stakeholders(doc: Document, items: list[Stakeholder]) -> None:
    _table(doc, ["Stakeholder", "Role", "Interest / Need", "Sources"],
           [[s.name, s.role, s.interest, _ids(s.source_ids)] for s in items])


def _risks(doc: Document, items: list[Risk]) -> None:
    _table(doc, ["#", "Risk", "Impact", "Mitigation", "Sources"],
           [[f"R-{i:02d}", r.description, r.impact, r.mitigation, _ids(r.source_ids)]
            for i, r in enumerate(items, start=1)])


def _conflicts(doc: Document, facts: ExtractedFacts) -> None:
    doc.add_paragraph("Sources contradict each other on these points. No option has been chosen; "
                      "a stakeholder decision is required.")
    if not facts.conflicts:
        _italic(doc, "No conflicts detected between sources.")
        return
    for c in facts.conflicts:
        doc.add_heading(f"{c.id} — {c.topic}", level=3)
        _table(doc, ["Option", "Statement", "Sources"],
               [[chr(ord("A") + i) if i < 26 else str(i + 1), o.statement, _ids(o.source_ids)]
                for i, o in enumerate(c.options)])
        meta = doc.add_paragraph()
        meta.add_run("Related requirements: ").bold = True
        meta.add_run(f"{_ids(c.related_requirement_ids)}    ")
        meta.add_run("Resolution: ").bold = True
        meta.add_run(c.resolution)


def _open_questions(doc: Document, facts: ExtractedFacts) -> None:
    doc.add_paragraph("These questions need a stakeholder answer before this document can be finalised.")
    _table(doc, ["Id", "Question", "Reason", "Related", "Priority", "Sources"],
           [[q.id, q.question, q.reason, _ids(q.related_requirement_ids), q.priority, _ids(q.source_ids)]
            for q in facts.open_questions])


def _review_notes(doc: Document, run: RunInfo) -> None:
    doc.add_paragraph("Automated validation and security checks found the following items to review:")
    if run.review_notes:
        _bullets(doc, run.review_notes)
    else:
        _italic(doc, "No issues found.")


def _change_log(doc: Document, run: RunInfo) -> None:
    labels = {"added": "Added requirements", "modified": "Modified requirements",
              "removed": "Removed requirements", "newly_answered": "Newly answered TBDs",
              "new_conflicts": "Newly discovered conflicts", "resolved_conflicts": "Resolved conflicts"}
    doc.add_paragraph(f"Version {run.version}.")
    _table(doc, ["Change", "Items"],
           [[label, ", ".join(run.change_log.get(key, [])) or "-"] for key, label in labels.items()])


def _sources_appendix(doc: Document, sources: list[SourceDoc]) -> None:
    doc.add_paragraph("Source ids used throughout this document refer to:")
    _table(doc, ["Id", "Type", "Title / Subject", "Date", "File"],
           [[s.id, s.kind, s.title, (s.date or "")[:10], Path(s.path).name] for s in sources])


def _by_type(facts: ExtractedFacts, kind: str) -> list[Requirement]:
    return [r for r in facts.requirements if r.type == kind]


# ---------------------------------------------------------------- BRD

def render_brd(brd: BRD, facts: ExtractedFacts, project: str, sources: list[SourceDoc],
               path: Path, run: RunInfo | None = None) -> None:
    run = run or RunInfo()
    doc = _new_document()
    _title_page(doc, "BRD", "Business Requirements Document", project, facts, len(sources), run)

    doc.add_heading("1. Executive Summary", level=1)
    _para(doc, brd.executive_summary)
    doc.add_heading("2. Background", level=1)
    _para(doc, brd.background)
    doc.add_heading("3. Business Problem", level=1)
    _para(doc, brd.problem_statement)

    doc.add_heading("4. Business Objectives", level=1)
    _table(doc, ["Id", "Objective", "Success Metric", "Sources"],
           [[o.id, o.statement, o.success_metric, _ids(o.source_ids)] for o in brd.business_objectives])

    doc.add_heading("5. Stakeholders", level=1)
    _stakeholders(doc, brd.stakeholders)

    doc.add_heading("6. Business Scope", level=1)
    doc.add_heading("6.1 In Scope", level=2)
    _bullets(doc, brd.scope_in)
    doc.add_heading("6.2 Out of Scope", level=2)
    _bullets(doc, brd.scope_out)

    doc.add_heading("7. Current Process", level=1)
    _para(doc, brd.current_process)
    doc.add_heading("8. Proposed Process", level=1)
    _para(doc, brd.proposed_process)

    doc.add_heading("9. Business Requirements", level=1)
    _legend(doc)
    _requirement_table(doc, _by_type(facts, "business"))

    doc.add_heading("10. Business Rules", level=1)
    _bullets(doc, brd.business_rules)
    doc.add_heading("11. Assumptions", level=1)
    _bullets(doc, brd.assumptions)
    doc.add_heading("12. Constraints", level=1)
    _bullets(doc, brd.constraints)
    doc.add_heading("13. Dependencies", level=1)
    _bullets(doc, brd.dependencies)
    doc.add_heading("14. Risks", level=1)
    _risks(doc, brd.risks)
    doc.add_heading("15. Success Criteria", level=1)
    _bullets(doc, brd.success_criteria)

    doc.add_heading("16. Conflicts", level=1)
    _conflicts(doc, facts)
    doc.add_heading("17. Open Questions", level=1)
    _open_questions(doc, facts)

    doc.add_heading("18. Review Notes", level=1)
    _review_notes(doc, run)
    doc.add_heading("Appendix A. Change Log", level=1)
    _change_log(doc, run)
    doc.add_heading("Appendix B. Sources", level=1)
    _sources_appendix(doc, sources)
    doc.save(str(path))


# ---------------------------------------------------------------- SRS

NFR_GROUPS = [
    ("Security Requirements", {"security", "compliance", "privacy"}),
    ("Performance Requirements", {"performance", "scalability", "capacity"}),
    ("Logging and Monitoring Requirements", {"logging", "monitoring", "auditing", "audit", "observability"}),
]


def render_srs(srs: SRS, facts: ExtractedFacts, project: str, sources: list[SourceDoc],
               path: Path, run: RunInfo | None = None) -> None:
    run = run or RunInfo()
    doc = _new_document()
    _title_page(doc, "SRS", "Software Requirements Specification", project, facts, len(sources), run)
    brs = _by_type(facts, "business")
    frs = _by_type(facts, "functional")
    nfrs = _by_type(facts, "non_functional")

    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Purpose", level=2)
    _para(doc, srs.purpose)
    doc.add_heading("1.2 Scope", level=2)
    _para(doc, srs.scope)
    doc.add_heading("1.3 Definitions, Acronyms and Abbreviations", level=2)
    _table(doc, ["Term", "Definition"], [[t.term, t.definition] for t in srs.definitions])

    doc.add_heading("2. Overall Description", level=1)
    doc.add_heading("2.1 System Overview", level=2)
    _para(doc, srs.system_overview)
    doc.add_heading("2.2 User Roles", level=2)
    _table(doc, ["User Role", "Description"], [[u.name, u.description] for u in srs.user_classes])
    doc.add_heading("2.3 Operating Environment", level=2)
    _para(doc, srs.operating_environment)
    doc.add_heading("2.4 Constraints", level=2)
    _bullets(doc, srs.design_constraints)
    doc.add_heading("2.5 Dependencies", level=2)
    _bullets(doc, srs.dependencies)
    doc.add_heading("2.6 Assumptions", level=2)
    _bullets(doc, srs.assumptions)

    doc.add_heading("3. System Workflows", level=1)
    for wf in srs.system_workflows:
        doc.add_heading(wf.name, level=2)
        for step in wf.steps:
            doc.add_paragraph(step, style="List Number")
        doc.add_paragraph(f"Requirements: {_ids(wf.requirement_ids)}")
    if not srs.system_workflows:
        _italic(doc, EMPTY)

    doc.add_heading("4. Functional Requirements", level=1)
    _legend(doc)
    _requirement_table(doc, frs, with_links=True)
    doc.add_heading("4.1 Acceptance Criteria", level=2)
    with_criteria = [r for r in frs + nfrs if r.acceptance_criteria]
    for r in with_criteria:
        doc.add_paragraph().add_run(f"{r.id} {r.title}").bold = True
        _bullets(doc, r.acceptance_criteria)
    if not with_criteria:
        _italic(doc, "TBD — the sources do not state acceptance criteria.")

    doc.add_heading("5. Non-Functional Requirements", level=1)
    _requirement_table(doc, nfrs, with_category=True, with_links=True)
    for i, (heading, categories) in enumerate(NFR_GROUPS, start=1):
        doc.add_heading(f"5.{i} {heading}", level=2)
        _bullets(doc, [f"{n.id} [{n.status}]: {n.requirement}" for n in nfrs
                       if n.category.strip().lower() in categories])

    doc.add_heading("6. Data Requirements", level=1)
    _table(doc, ["Entity", "Description", "Key Attributes", "Sources"],
           [[d.name, d.description, ", ".join(d.key_attributes) or "TBD", _ids(d.source_ids)]
            for d in srs.data_requirements])

    doc.add_heading("7. Integration and External Interfaces", level=1)
    _table(doc, ["Interface", "Type", "Description", "Sources"],
           [[i.name, i.type, i.description, _ids(i.source_ids)] for i in srs.external_interfaces])

    doc.add_heading("8. Error Handling", level=1)
    _bullets(doc, srs.error_handling)

    doc.add_heading("9. Conflicts", level=1)
    _conflicts(doc, facts)
    doc.add_heading("10. Open Questions", level=1)
    _open_questions(doc, facts)

    doc.add_heading("11. Requirements Traceability", level=1)
    doc.add_heading("11.1 Requirements Traceability Matrix", level=2)
    _table(doc, ["Requirement", "Type", "Source", "Source reference", "Status", "Confidence"],
           [[r.id, r.type.replace("_", "-"), _ids(r.source_ids), r.source_reference or "-",
             r.status, r.confidence] for r in facts.requirements],
           [_fill(r.status) for r in facts.requirements])
    doc.add_heading("11.2 Business → Technical Requirements", level=2)
    doc.add_paragraph("Business requirements mapped to the functional and non-functional "
                      "requirements that serve them.")
    by_br: dict[str, list[str]] = {br.id: [] for br in brs}
    for r in frs + nfrs:
        for ref in r.br_refs:
            by_br.setdefault(ref, []).append(r.id)
    titles = {br.id: br.title for br in brs}
    _table(doc, ["BR", "Business Requirement", "Implemented by"],
           [[br_id, titles.get(br_id, "(not in registry)"), ", ".join(ids) or "NOT COVERED"]
            for br_id, ids in by_br.items()])

    doc.add_heading("12. Review Notes", level=1)
    _review_notes(doc, run)
    doc.add_heading("Appendix A. Change Log", level=1)
    _change_log(doc, run)
    doc.add_heading("Appendix B. Sources", level=1)
    _sources_appendix(doc, sources)
    doc.save(str(path))
