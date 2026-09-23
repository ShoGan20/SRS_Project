"""Text extraction for meeting notes: .txt, .md, .docx, .pdf, .vtt."""

from __future__ import annotations

import re
from pathlib import Path


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(f"[Page {number}]\n{text}")
    return "\n\n".join(pages)


_VTT_TIMING = re.compile(r"^\s*(\d{1,2}:)?\d{2}:\d{2}[.,]\d{3}\s*-->")
_VTT_VOICE = re.compile(r"<v(?:\.[^ >]*)?\s+([^>]+)>(.*?)(?:</v>|$)")
_VTT_TAG = re.compile(r"<[^>]+>")


def read_vtt(path: Path) -> str:
    """Flatten a WebVTT transcript to 'Speaker: text' lines, merging consecutive
    cues from the same speaker (Teams/Zoom split sentences across many cues)."""
    lines = read_text_file(path).splitlines()
    turns: list[list[str]] = []  # [speaker, text]
    in_cue = False
    for raw in lines:
        line = raw.strip()
        if _VTT_TIMING.match(line):
            in_cue = True
            continue
        if not line:
            in_cue = False
            continue
        if not in_cue:
            continue  # header, NOTE blocks, cue ids
        speaker = ""
        match = _VTT_VOICE.search(line)
        if match:
            speaker, line = match.group(1).strip(), match.group(2)
        elif ":" in line and len(line.split(":", 1)[0]) <= 40:
            # Zoom style "Name: text"
            speaker, line = (s.strip() for s in line.split(":", 1))
        text = _VTT_TAG.sub("", line).strip()
        if not text:
            continue
        if turns and turns[-1][0] == speaker:
            turns[-1][1] += " " + text
        else:
            turns.append([speaker, text])
    return "\n".join(f"{s}: {t}" if s else t for s, t in turns)


READERS = {
    ".txt": read_text_file,
    ".md": read_text_file,
    ".docx": read_docx,
    ".pdf": read_pdf,
    ".vtt": read_vtt,
}
