"""Generate the binary sample notes (.docx and .pdf) next to the text ones.

    python samples/make_samples.py
"""

from pathlib import Path

from docx import Document

MEETINGS = Path(__file__).parent / "meetings"


def make_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Phoenix - Steering Group", level=1)
    doc.add_paragraph("Date: 30 June 2026")
    doc.add_paragraph("Attendees: Priya Nair, Tom Becker, Aisha Khan, Marco Rossi, Helen Ward (Finance Director)")
    doc.add_heading("Decisions", level=2)
    for text in [
        "Go-live target confirmed for end of November 2026.",
        "Customers must be able to update their contact details (email, phone, postal address) in the portal.",
        "Payment method scope (card only vs. card + direct debit at launch) NOT resolved - Helen and Priya to agree by 10 July.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Budget", level=2)
    table = doc.add_table(rows=3, cols=2)
    for row, (item, value) in zip(table.rows, [("Build (one-off)", "GBP 1.2m"),
                                               ("Run (per year)", "GBP 180k"),
                                               ("Contingency", "10%")]):
        row.cells[0].text, row.cells[1].text = item, value
    doc.save(str(path))


def make_pdf(path: Path, lines: list[str]) -> None:
    """Write a one-page PDF with plain Helvetica text (no extra libraries needed)."""
    def esc(s: str) -> str:
        return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    stream = "BT /F1 11 Tf 50 790 Td 15 TL\n" + "\n".join(f"({esc(l)}) '" for l in lines) + "\nET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
        "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(stream.encode('latin-1'))} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("latin-1")
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode("latin-1")
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("latin-1")
    path.write_bytes(out)


if __name__ == "__main__":
    MEETINGS.mkdir(exist_ok=True)
    make_docx(MEETINGS / "2026-06-30_steering_group.docx")
    make_pdf(MEETINGS / "2026-07-07_ux_review.pdf", [
        "Phoenix - UX Review",
        "Date: 7 July 2026   Attendees: Priya Nair, Sam Okoro (UX Lead), Aisha Khan",
        "",
        "- Usability testing with 12 customers (6 aged 65+).",
        "- Registration must take under 5 minutes for 90% of test users.",
        "- Bill history: show a usage chart comparing the last 12 months.",
        "- Meter reading: allow uploading a photo of the meter as evidence.",
        "- Support English and Welsh languages at launch.",
        "- Open: do we need a live chat widget? No budget identified yet.",
    ])
    print("Sample .docx and .pdf written to", MEETINGS)
